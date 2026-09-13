#!/usr/bin/env python3
"""Create the .xlsx twin a draft TAPP CSV needs, for the drafts that lack one.

build_tapp_examples reads publication columns through openpyxl and refuses a bare .csv:

    <name>.csv is a .csv and no .xlsx twin exists beside it; this builder needs the
    workbook form to read publication columns

30 of the 43 draft TAPPs ship a twin, 13 do not, and those 13 are exactly the techniques
whose `examples` stage fails in a full regeneration. This writes the missing ones.

A twin carries two sheets, matching the 30 that already exist:

  TAPP     the CSV verbatim -- same header, same rows, same order. Empty CSV cells become
           empty cells, not the string "".
  Legends  the tier and mode-column definitions. These are NOT in the CSV. They are shared
           boilerplate: 29 of the 30 existing twins carry byte-identical Legends (VNMIR has
           a 32-row variant), so the common 28-row block is copied from a donor twin rather
           than invented. Nothing here is technique-specific.

Only the TAPP sheet is read by the builder; Legends is carried so a generated twin is not
structurally poorer than a delivered one.

    python tools/make_draft_xlsx_twin.py            # report what is missing
    python tools/make_draft_xlsx_twin.py --write
"""
import argparse
import csv
import glob
import os
import sys

import openpyxl

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
DRAFTS = os.path.join(ROOT, "draftTAPPs")
DONOR = os.path.join(DRAFTS, "AIVA_TAPP_draft_v2.xlsx")


def missing():
    out = []
    for c in sorted(glob.glob(os.path.join(DRAFTS, "*.csv"))):
        if not os.path.exists(os.path.splitext(c)[0] + ".xlsx"):
            out.append(c)
    return out


def legends_from(donor):
    wb = openpyxl.load_workbook(donor, read_only=True, data_only=True)
    if "Legends" not in wb.sheetnames:
        raise SystemExit("donor %s has no Legends sheet" % donor)
    return [list(r) for r in wb["Legends"].iter_rows(values_only=True)]


def build(csv_path, legends, write):
    rows = list(csv.reader(open(csv_path, encoding="utf-8-sig")))
    dest = os.path.splitext(csv_path)[0] + ".xlsx"
    if not write:
        return dest, len(rows) - 1, len(rows[0]) if rows else 0
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "TAPP"
    for r in rows:
        # a CSV has no null, only "", but the delivered twins leave those cells truly empty;
        # writing None keeps a generated twin indistinguishable from a delivered one
        ws.append([c if c != "" else None for c in r])
    lg = wb.create_sheet("Legends")
    for r in legends:
        lg.append(list(r))
    wb.save(dest)
    return dest, len(rows) - 1, len(rows[0]) if rows else 0


def verify(csv_path):
    """A twin is good only if it round-trips: same header, same row count, same cells."""
    dest = os.path.splitext(csv_path)[0] + ".xlsx"
    rows = list(csv.reader(open(csv_path, encoding="utf-8-sig")))
    wb = openpyxl.load_workbook(dest, read_only=True, data_only=True)
    if wb.sheetnames != ["TAPP", "Legends"]:
        return "sheets are %s" % wb.sheetnames
    got = [["" if c is None else str(c) for c in r]
           for r in wb["TAPP"].iter_rows(values_only=True)]
    if len(got) != len(rows):
        return "row count %d != %d" % (len(got), len(rows))
    for i, (a, b) in enumerate(zip(rows, got)):
        a = a + [""] * (len(b) - len(a))
        if a != b[:len(a)]:
            return "row %d differs" % i
    lit = [j for j, h in enumerate(got[0]) if h.strip().lower() == "literature assessment"]
    if not lit:
        return "no 'Literature Assessment' column -- the builder would find no publications"
    pub = [h for h in got[0][lit[0] + 1:] if h.strip()]
    return None if pub else "no publication columns after 'Literature Assessment'"


def main():
    ap = argparse.ArgumentParser(description=__doc__,
                                 formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--write", action="store_true")
    ap.add_argument("--donor", default=DONOR, help="twin to copy the Legends sheet from")
    a = ap.parse_args()

    todo = missing()
    if not todo:
        print("every draft TAPP csv already has an .xlsx twin")
        return 0
    legends = legends_from(a.donor) if a.write else []
    if a.write:
        print("Legends copied from %s (%d rows)\n" % (os.path.basename(a.donor), len(legends)))
    bad = 0
    for c in todo:
        dest, nrows, ncols = build(c, legends, a.write)
        note = ""
        if a.write:
            err = verify(c)
            if err:
                note = "  VERIFY FAILED: " + err
                bad += 1
        print("%-42s %3d rows x %2d cols -> %s%s"
              % (os.path.basename(c), nrows, ncols, os.path.basename(dest), note))
    print("\n%d twin(s) %s" % (len(todo), "written" if a.write else "would be written"))
    if a.write:
        print("%d verified, %d failed" % (len(todo) - bad, bad))
    return 1 if bad else 0


if __name__ == "__main__":
    sys.exit(main())
